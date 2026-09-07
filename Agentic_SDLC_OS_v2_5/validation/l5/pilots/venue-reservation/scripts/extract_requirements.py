#!/usr/bin/env python3
from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source" / "requirements.xlsx"
OUTPUT = ROOT / "source" / "requirements-selection.md"

SHEETS = {
    "details": "03_\u6838\u5fc3\u7279\u8272\u529f\u80fd\u8be6\u8bbe",
    "states": "04_\u72b6\u6001\u673a\u4e0e\u89c4\u5219\u5f15\u64ce",
    "metrics": "06_\u6570\u636e\u6307\u6807\u53e3\u5f84",
}
DETAIL_IDS = {"D26", "D33"}
STATE_IDS = {"Reservation", "WaitlistEntry", "SeatHold", "Session", "R02", "R08"}
METRIC_NAMES = {
    "\u9884\u7ea6\u6210\u529f\u7387",
    "\u5019\u8865\u8f6c\u6b63\u7387",
    "\u7a7a\u4f4d\u586b\u8865\u7387",
    "\u5e73\u5747\u5019\u8865\u7b49\u5f85\u65f6\u957f",
    "\u7403\u53f0\u5229\u7528\u7387",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def markdown_table(headers, rows) -> list[str]:
    safe = lambda value: str(value or "").replace("|", "\\|").replace("\n", "<br>")
    lines = ["| " + " | ".join(map(safe, headers)) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    lines.extend("| " + " | ".join(map(safe, row)) + " |" for row in rows)
    return lines


def selected_rows(sheet, first_column_values: set[str]):
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    selected = [row for row in rows[1:] if str(row[0]) in first_column_values]
    return headers, selected


def main() -> int:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    details_h, details = selected_rows(workbook[SHEETS["details"]], DETAIL_IDS)
    states_h, states = selected_rows(workbook[SHEETS["states"]], STATE_IDS)
    metrics_h, metrics = selected_rows(workbook[SHEETS["metrics"]], METRIC_NAMES)
    lines = [
        "# Venue Reservation Pilot - Source Requirements",
        "",
        f"- Source file: `requirements.xlsx`",
        f"- Source SHA-256: `{sha256(SOURCE)}`",
        "- Selection rule: D26, D33, reservation/waitlist/hold state invariants, and outcome metrics.",
        "- Purpose: preserve a reproducible, narrow vertical slice rather than reinterpret the full product backlog.",
        "",
        "## Product capabilities",
        "",
        *markdown_table(details_h, details),
        "",
        "## State machines and invariants",
        "",
        *markdown_table(states_h, states),
        "",
        "## Product metrics",
        "",
        *markdown_table(metrics_h, metrics),
        "",
        "## Pilot commitment",
        "",
        "The implementation must allocate capacity atomically, maintain one active reservation or waitlist entry per member and session, use FIFO promotion, make retries idempotent, preserve an immutable audit trail, expose actionable error states, and prove that concurrent demand cannot oversell capacity.",
    ]
    OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
